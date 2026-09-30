from __future__ import annotations
import re
import zipfile
from hashlib import sha256
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

path=Path(r'D:\高效管理\提交材料\项目概要介绍.docx')
if not path.is_file(): raise SystemExit('DOCX_MISSING')
with zipfile.ZipFile(path) as z:
    if z.testzip(): raise SystemExit('DOCX_ZIP_CORRUPT')
    xml=z.read('word/document.xml').decode('utf-8')
    sect=z.read('word/document.xml')

doc=Document(path)
required=['项目介绍','核心功能','技术路线','预期成果','创新性说明','应用价值说明']
all_text=[]
for p in doc.paragraphs: all_text.append(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            all_text.append(cell.text)
text='\n'.join(all_text)
for item in required:
    if item not in text: raise SystemExit(f'MISSING_SECTION={item}')
# Match the exact six body paragraphs by unique section openings.
starts=[
 '高校校园服务并不缺系统',
 '系统以东方神兽“白泽”',
 '项目采用“体验层',
 '形成可本地运行',
 '一是“白泽总调度',
 '项目直接回应高校报修',
]
body=[]
for start in starts:
    found=[p.text for p in doc.paragraphs if p.text.startswith(start)]
    if len(found)!=1: raise SystemExit(f'BODY_MATCH_FAILURE={start}:{len(found)}')
    body.append(found[0])
count=len(re.sub(r'\s','', ''.join(body)))
section=doc.sections[0]
page_w=round(section.page_width.cm,2); page_h=round(section.page_height.cm,2)
if abs(page_w-21.0)>0.1 or abs(page_h-29.7)>0.1: raise SystemExit(f'NOT_A4={page_w}x{page_h}')
if count>2000: raise SystemExit(f'OVER_LIMIT={count}')
print(f'DOCX_FINAL_VERIFY=PASS')
print(f'BODY_CHAR_COUNT_NO_SPACE={count}')
print(f'REQUIRED_SECTIONS={len(required)}')
print(f'PAGE_SIZE_CM={page_w}x{page_h}')
print(f'DOCX_BYTES={path.stat().st_size}')
print(f'SHA256={sha256(path.read_bytes()).hexdigest()}')
