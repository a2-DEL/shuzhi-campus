from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(r"D:\高效管理")
OUT_DIR = ROOT / "提交材料"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "项目概要介绍.docx"

TITLE = "数智星图——高校空间微治理白泽 Agent OS"
SUBTITLE = "项目概要介绍"
TAGLINE = "让 AI 完成工作，而不是只回答问题"

SECTIONS = [
    (
        "01",
        "项目介绍",
        "高校校园服务并不缺系统，真正缺少的是跨角色、跨部门、跨空间的责任识别与执行闭环。项目前期已完成500份问卷分析、156条报修位置记录整理、102条脱敏报修提交分析、43条失物状态追踪、86条测试反馈、四轮原型试错及竞品研究。调研显示，70%的受访者寻找空教室超过15分钟，60%的报修等待超过2天，信息分散和责任难定位是共同痛点。团队据此构建“数智星图——高校空间微治理白泽 Agent OS”，面向师生、院系、后勤、宿管和信息化部门，把人员、空间、设备、制度与任务连接为可执行校园星图。项目契合赛题对真实教育痛点、端到端智能体、知识库、工作流与工具集成的要求，目标是形成校园数字员工组织。",
    ),
    (
        "02",
        "核心功能",
        "系统以东方神兽“白泽”为统一AI助手和总调度官。用户提出目标后，白泽识别问答、参谋或执行模式，读取权限与业务上下文，拆解任务、生成预览并等待人工裁决；获批后，将任务分发给玄龟（后勤执行）、灵鹊（消息触达）、獬豸（合规校验）等专业Agent，在全屏星图中呈现接令、交接、工具调用、数据库写回、回读验证和成果归整。系统已形成14类角色团队、45个专业Agent、8个核心Skill，覆盖报修与SLA、教室预约、通知触达、宿舍安全、卫生整改、访客准入、失物招领、能耗维护等8类业务闭环；并提供可视化工作流、Skills/插件/MCP、可信知识库与混合检索、知识图谱与GraphRAG、多Agent协同、受控自进化、数字孪生、多模态证据和AI运维中心。界面全部使用业务语言，不展示代码或JSON。",
    ),
    (
        "03",
        "技术路线",
        "项目采用“体验层—白泽总调度层—Agent OS运行层—认知与工具层—业务与治理层”分层架构。前端基于Next.js、React、TypeScript构建企业级工作台；后端以PostgreSQL作为业务与运行事实源，通过租户隔离、角色/属性权限、持久任务、事务Outbox、幂等与审计保障可靠性。白泽经受控模型网关调用DeepSeek，模型只负责意图理解、证据化回答和自然语言汇总，不能绕过权限改库；写操作必须经历Prepare—Preview—Approve—Revalidate—Commit—Verify。认知侧采用版本化知识库、权限裁剪、全文/向量混合检索、GraphRAG与引用校验；执行侧通过不可变Skill契约、可视化DAG、插件和MCP完成工具发现与调用；SSE订阅PostgreSQL持久事件，服务重启后可恢复。AI运维、模型台账、失败回滚、评测门禁、孪生沙盒和签名联邦贡献形成治理链。项目按R1可信业务版、R2 Agent OS企业版、R3认知进化版、R4多校前沿版推进，各阶段可独立验收。",
    ),
    (
        "04",
        "预期成果",
        "形成可本地运行、现场演示和持续扩展的高校空间微治理智能体系统及参赛成果包。系统侧交付白泽助手、多Agent运行时、8类业务闭环、知识图谱、工作流、Skill/MCP、AI运维、数字孪生和受控联邦；材料侧形成需求分析、智能体设计说明书、测试方案与结果、项目PPT和3—5分钟真实系统演示视频。后续在合作学校授权和基础设施具备后，逐步接入校级身份、IoT设备与跨校节点，开展真实环境试点和效能评估。",
    ),
    (
        "05",
        "创新性说明",
        "一是“白泽总调度+分灵体团队”组织创新，把单一问答机器人升级为有岗位、权限、Skill、审批和绩效证据的数字员工团队，并以真实事件驱动星图展示协同。二是“业务结果反向验收”：Agent完成工具调用不等于任务成功，必须经业务表写入、数据库回读、审计账本和业务效果交叉验证后，白泽才可汇报，解决智能体“说成功、未落地”的难题。三是“知识—工具—业务—治理”一体化：可追溯GraphRAG负责认知，Skill/MCP负责执行，人工裁决、确定性校验和事务链负责安全；自进化模块从失败与反馈生成候选策略，经固定数据集、安全门禁、人工发布和一键回滚后才能运行。四是以校园空间本体连接人员、楼宇、设备、制度、工单和Agent，为数字孪生推演与多校隐私协作提供统一语义底座。",
    ),
    (
        "06",
        "应用价值说明",
        "项目直接回应高校报修、预约、通知、访客、宿舍与能耗等高频痛点，把“查询—转发—催办—统计”的重复劳动转化为可审计的Agent协作链，使事项自动找到正确责任人、工具和处理路径。当前本地脱敏环境已通过真实浏览器动态验收：三Agent任务完成3项业务效果验证；知识中枢形成10份文档、30个实体、15条发布关系；MCP可发现并调用3项工具；数字孪生推演和3节点签名聚合可真实运行，关键页面无运行异常。系统采用模块化、多租户设计，可从院系扩展到校级并复制到不同高校。其价值不仅是提升响应与管理透明度，更在于沉淀可复用的校园知识、责任图谱、业务Skill和治理规范，为高校“人工智能+教育管理”提供可持续演进的工程样板。",
    ),
]
FOOTNOTES = [
    "数据口径：500份问卷及70%/60%为调研样本统计；156条为位置记录，102条为脱敏报修提交，86条为测试反馈。",
    "当前成果运行于本地脱敏演示环境；跨校生产联邦、真实IoT和校级部署属于后续试点阶段，不作既成成果宣称。",
]

KEY_METRICS = [
    ("500", "份问卷分析"),
    ("8", "类业务闭环"),
    ("14", "类角色团队"),
    ("45", "个专业 Agent"),
    ("8", "个核心 Skill"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header_noop(row) -> None:
    set_repeat_table_header(row)


def set_table_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="25324A") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.45) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = True
    fmt.widow_control = True


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, fld_char_end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("数智星图 · 项目概要介绍   ")
    set_font(r, size=8.5, color="66748A")
    add_field(paragraph, "PAGE")


def count_text() -> tuple[str, int, int]:
    content = "".join(text for _, _, text in SECTIONS)
    count_no_space = len(re.sub(r"\s", "", content))
    count_chinese = len(re.findall(r"[\u4e00-\u9fff]", content))
    return content, count_no_space, count_chinese


def add_title_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.different_first_page_header_footer = True
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # Brand eyebrow
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=5, after=7, line=1)
    r = p.add_run("BAIZE AGENT OS  ·  AI AGENT COMPETITION SUBMISSION")
    set_font(r, name="Aptos", size=9, bold=True, color="18A5B8")

    # Accent line
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    widths = [1.0, 12.8, 1.0]
    colors = ["10B9C8", "0D1B35", "E3B85C"]
    for i, cell in enumerate(t.rows[0].cells):
        set_table_cell_width(cell, int(widths[i] * 567))
        set_cell_shading(cell, colors[i])
        cell.text = ""
        cell.height = Cm(0.12)
        set_cell_margins(cell, 0, 0, 0, 0)
        set_cell_border(cell, top={"val": "nil"}, left={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=17, after=4, line=1.1)
    r = p.add_run(TITLE)
    set_font(r, size=24, bold=True, color="0B1832")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=1, after=6, line=1)
    r = p.add_run(SUBTITLE)
    set_font(r, size=15, bold=True, color="18A5B8")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=14, line=1.2)
    r = p.add_run(TAGLINE)
    set_font(r, size=11.5, color="6B778C")

    # Positioning card
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(16.8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F8FA")
    set_cell_margins(cell, top=170, start=260, bottom=170, end=260)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 8, "color": "A6DDE3"},
        left={"val": "single", "sz": 8, "color": "A6DDE3"},
        bottom={"val": "single", "sz": 8, "color": "A6DDE3"},
        right={"val": "single", "sz": 8, "color": "A6DDE3"},
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目定位  |  高校空间微治理多智能体操作系统")
    set_font(r, size=11, bold=True, color="0B6E7D")
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=3, after=0, line=1.3)
    r = p.add_run("真实痛点驱动 · 多 Agent 协同 · 业务闭环验证 · 企业级安全治理")
    set_font(r, size=9.5, color="466176")

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=15, after=5, line=1)
    r = p.add_run("项目证据与落地基线")
    set_font(r, size=11.5, bold=True, color="0B1832")

    metrics = doc.add_table(rows=1, cols=len(KEY_METRICS))
    metrics.autofit = False
    for i, (number, label) in enumerate(KEY_METRICS):
        cell = metrics.cell(0, i)
        set_cell_shading(cell, "0D1B35" if i % 2 == 0 else "102743")
        set_cell_margins(cell, top=120, start=70, bottom=110, end=70)
        set_cell_border(
            cell,
            top={"val": "single", "sz": 4, "color": "203A5B"},
            left={"val": "single", "sz": 4, "color": "203A5B"},
            bottom={"val": "single", "sz": 4, "color": "203A5B"},
            right={"val": "single", "sz": 4, "color": "203A5B"},
        )
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=1, line=1)
        r = p.add_run(number)
        set_font(r, name="Aptos Display", size=20, bold=True, color="65E0E8" if i < 3 else "EAC66D")
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, line=1)
        r = p.add_run(label)
        set_font(r, size=8.5, color="EAF3FA")

    # Scoring adaptation grid
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=14, after=5, line=1)
    r = p.add_run("与赛题评分的直接适配")
    set_font(r, size=11.5, bold=True, color="0B1832")

    adapt = doc.add_table(rows=2, cols=3)
    adapt.autofit = False
    adapts = [
        ("方案完整性", "需求—设计—开发—测试—演示全链路"),
        ("技术创新性", "白泽总调度、GraphRAG、Skill/MCP、受控进化"),
        ("场景应用价值", "8类校园业务闭环与数据库回读验证"),
        ("路演可演示性", "全屏Agent星图、真实等待、事件与审计证据"),
        ("作品优化效果", "从业务平台升级为企业级Agent OS"),
        ("可推广潜力", "多角色、多租户、模块化与R1—R4路线"),
    ]
    for idx, (head, body) in enumerate(adapts):
        cell = adapt.cell(idx // 3, idx % 3)
        set_cell_shading(cell, "F7F9FC")
        set_cell_margins(cell, top=100, start=130, bottom=100, end=130)
        set_cell_border(
            cell,
            top={"val": "single", "sz": 5, "color": "D5DEE9"},
            left={"val": "single", "sz": 5, "color": "D5DEE9"},
            bottom={"val": "single", "sz": 5, "color": "D5DEE9"},
            right={"val": "single", "sz": 5, "color": "D5DEE9"},
        )
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=2, line=1)
        r = p.add_run(head)
        set_font(r, size=9.5, bold=True, color="0B6E7D")
        p = cell.add_paragraph()
        set_paragraph_spacing(p, line=1.25)
        r = p.add_run(body)
        set_font(r, size=8.6, color="405066")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=13, after=0, line=1)
    r = p.add_run("申报方向：校园管理类 AI 智能体  |  文档版本：V1.0  |  2026年8月")
    set_font(r, size=8.5, color="8793A6")


def add_content_page(doc: Document, count_no_space: int) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4, line=1)
    r = p.add_run("PROJECT OVERVIEW  /  核心概要")
    set_font(r, name="Aptos", size=9, bold=True, color="18A5B8")
    r = p.add_run(f"                                      正文字数：{count_no_space}（不含封面及口径说明）")
    set_font(r, size=8.5, color="7A8798")

    for number, heading, text in SECTIONS:
        if number == "04":
            doc.add_page_break()
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        left, right = table.cell(0, 0), table.cell(0, 1)
        set_table_cell_width(left, 1000)
        set_table_cell_width(right, 8200)
        set_cell_shading(left, "0D1B35")
        set_cell_shading(right, "EAF7F9")
        for cell in (left, right):
            set_cell_margins(cell, top=70, start=110, bottom=70, end=110)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 4, "color": "B9DDE2"},
                left={"val": "single", "sz": 4, "color": "B9DDE2"},
                bottom={"val": "single", "sz": 4, "color": "B9DDE2"},
                right={"val": "single", "sz": 4, "color": "B9DDE2"},
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = left.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p1, line=1)
        r = p1.add_run(number)
        set_font(r, name="Aptos Display", size=12, bold=True, color="65E0E8")
        p2 = right.paragraphs[0]
        set_paragraph_spacing(p2, line=1)
        r = p2.add_run(heading)
        set_font(r, size=11.5, bold=True, color="0B6E7D")

        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=3, after=7, line=1.43)
        p.paragraph_format.first_line_indent = Pt(21)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        set_font(r, size=10.25, color="28364A")

    # Evidence boundary box
    box = doc.add_table(rows=1, cols=1)
    cell = box.cell(0, 0)
    set_cell_shading(cell, "FFF9EC")
    set_cell_margins(cell, top=90, start=150, bottom=80, end=150)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 6, "color": "E7C66C"},
        left={"val": "single", "sz": 6, "color": "E7C66C"},
        bottom={"val": "single", "sz": 6, "color": "E7C66C"},
        right={"val": "single", "sz": 6, "color": "E7C66C"},
    )
    p = cell.paragraphs[0]
    r = p.add_run("真实性与成果边界")
    set_font(r, size=9.2, bold=True, color="9A6D00")
    for note in FOOTNOTES:
        p = cell.add_paragraph(style=None)
        set_paragraph_spacing(p, before=1, after=0, line=1.25)
        p.paragraph_format.left_indent = Pt(3)
        r = p.add_run("• " + note)
        set_font(r, size=8.2, color="66572D")


def add_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = SUBTITLE
    props.subject = "数智星图高校空间微治理白泽 Agent OS 参赛项目概要"
    props.author = "数智星图项目团队"
    props.keywords = "AI Agent, 白泽, 高校空间微治理, 多智能体, GraphRAG, MCP, 数字孪生"
    props.comments = "依据真实调研、系统实现和验收证据编制；正文字数不超过2000字。"


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("数智星图  BAIZE AGENT OS")
        set_font(r, name="Aptos", size=8, bold=True, color="7F8CA0")
        footer = section.footer
        p = footer.paragraphs[0]
        add_page_number(p)


def set_document_defaults(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = RGBColor.from_string("28364A")
    normal.paragraph_format.line_spacing = 1.43
    normal.paragraph_format.space_after = Pt(0)


def validate_docx(path: Path, expected_text: str, count_no_space: int) -> None:
    if not path.exists() or path.stat().st_size < 20_000:
        raise RuntimeError(f"DOCX missing or unexpectedly small: {path}")
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Corrupt member: {bad}")
        xml = zf.read("word/document.xml").decode("utf-8")
        for required in [TITLE, "项目介绍", "核心功能", "技术路线", "预期成果", "创新性说明", "应用价值说明"]:
            if required not in xml:
                raise RuntimeError(f"Missing required text: {required}")
    loaded = Document(path)
    all_text = "\n".join(p.text for p in loaded.paragraphs)
    # Table content is checked in XML; body section texts must survive paragraph round-trip.
    for fragment in ["500份问卷分析", "东方神兽“白泽”为统一AI助手和总调度官", "Prepare—Preview—Approve—Revalidate—Commit—Verify", "三Agent任务完成3项业务效果验证"]:
        if fragment not in all_text:
            raise RuntimeError(f"Missing body fragment after save: {fragment}")
    if count_no_space > 2000:
        raise RuntimeError(f"正文超过2000字：{count_no_space}")


def main() -> None:
    _, count_no_space, count_chinese = count_text()
    print(f"BODY_CHAR_COUNT_NO_SPACE={count_no_space}")
    print(f"BODY_CHINESE_CHAR_COUNT={count_chinese}")
    if count_no_space > 2000:
        raise SystemExit("正文超过2000字，请压缩后再生成")

    doc = Document()
    set_document_defaults(doc)
    add_document_properties(doc)
    add_title_page(doc)
    add_content_page(doc, count_no_space)
    add_header_footer(doc)
    doc.save(OUT)
    validate_docx(OUT, "".join(t for _, _, t in SECTIONS), count_no_space)
    print(f"DOCX_CREATED={OUT}")
    print(f"DOCX_BYTES={OUT.stat().st_size}")
    print("DOCX_VALIDATION=PASS")


if __name__ == "__main__":
    main()




